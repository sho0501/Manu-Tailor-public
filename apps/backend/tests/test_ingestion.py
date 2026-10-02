from io import BytesIO
import json

from docx import Document as WordDocument
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from fastapi.testclient import TestClient

from app.ingestion import ingest
from app.main import app
from app.db import connect


def text_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 750 Td (1. Set load to 10kg.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_text_layer_without_ocr(tmp_path, monkeypatch):
    path = tmp_path / "source.pdf"
    path.write_bytes(text_pdf())

    def forbidden(_):
        raise AssertionError("OCR must not be used for a text PDF")

    monkeypatch.setattr("app.ingestion.ocr", forbidden)
    document = ingest(path, "PDF sample")
    assert "10kg" in document.blocks[0].source_text
    assert document.blocks[0].source_page == 1
    assert "10kg" in document.raw_pages[0].text


def test_import_saves_untouched_raw_version_before_organization():
    original = "お願い\n・傷を付けないでください。\n\n条件が合う場合は、設定を確認します。"
    with TestClient(app) as client:
        login = client.post("/api/login", json={"username": "admin", "password": "manutailor-demo"}).json()
        headers = {"Authorization": "Bearer " + login["token"]}
        response = client.post("/api/manuals", headers=headers,
                               data={"title": "原文保存", "text": original})
        assert response.status_code == 200, response.text
        manual_id = response.json()["id"]
        with connect() as db:
            versions = db.execute("SELECT document FROM manual_versions WHERE manual_id=? ORDER BY version",
                                  (manual_id,)).fetchall()
        raw, organized = [json.loads(row["document"]) for row in versions]
        assert raw["organization_status"] == "raw"
        assert raw["blocks"] == []
        assert raw["raw_pages"] == [{"page": 1, "text": original}]
        assert organized["raw_pages"] == raw["raw_pages"]
        assert organized["blocks"]


def test_image_is_retained_without_ocr(tmp_path, monkeypatch):
    path = tmp_path / "diagram.png"
    Image.new("RGB", (200, 100), "white").save(path)
    monkeypatch.setattr("app.ingestion.ocr", lambda _: "")
    document = ingest(path, "図")
    assert len(document.images) == 1
    assert document.images[0].description is None
    assert document.extraction_notes


def test_pdf_embedded_figure_is_extracted_with_page_tags(tmp_path):
    path = tmp_path / "diagram.pdf"
    Image.new("RGB", (200, 100), "blue").save(path, format="PDF")
    document = ingest(path, "図入りPDF")
    assert any(image.tags == ["図", "ページ1"] for image in document.images)
    assert all(image.page == 1 for image in document.images)


def test_docx_table_and_markdown_heading(tmp_path):
    word = WordDocument()
    word.add_paragraph("1. 原稿を置いてください。")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "上限"
    table.cell(0, 1).text = "10kg"
    path = tmp_path / "source.docx"
    word.save(str(path))
    assert any("10kg" in b.source_text for b in ingest(path, "Word").blocks)
    path = tmp_path / "source.md"
    path.write_text("# 開始\n1. 電源を確認。\n注意：電源を切らないでください。", encoding="utf-8")
    document = ingest(path, "Markdown")
    assert document.blocks[0].heading == "開始"
    assert document.blocks[0].kind != "warning"
    assert document.blocks[1].kind == "warning"


def test_wrapped_pdf_text_and_requests_are_not_fragmented_steps():
    from app.ingestion import structure

    text = """キャ ビネ ッ ト・とびら
かたくしぼった、 ぬれ布巾で拭く。
・  落ちにくいよごれは、 ぬれた布巾をよごれの上に置いて30分
ぐらいふやかしてから拭きます。よごれがひどい場合は薄め
た台所用洗剤 （中性） をつけた布で拭き取り、 ぬれた布巾で洗
剤をよく拭き取ってください。
とびらの内側・庫内
お願い
・  たわしやフォークなど先のとがった物でこすらないで
ください。傷付いたり、 割れる原因になります。
分類できない説明文"""
    document = structure([(46, text)], "電子レンジ", merge_wrapped=True)
    values = [block.source_text for block in document.blocks]
    assert any("30分ぐらいふやかしてから拭きます。" in value for value in values)
    assert any("薄めた台所用洗剤" in value for value in values)
    assert not any(value.startswith("ぐらい") or value.startswith("た台所") for value in values)
    assert next(block for block in document.blocks if block.source_text == "お願い").kind == "request"
    assert next(block for block in document.blocks if "たわしやフォーク" in block.source_text).kind == "request"
    assert next(block for block in document.blocks if block.source_text == "分類できない説明文").kind == "other"


def test_restructure_and_custom_tags_create_versions():
    with TestClient(app) as client:
        login = client.post("/api/login", json={"username": "admin", "password": "manutailor-demo"}).json()
        headers = {"Authorization": "Bearer " + login["token"]}
        created = client.post("/api/manuals", headers=headers, data={
            "title": "分類検証", "text": "手入れ\n・ 汚れを30分\n置いて拭きます。\nお願い\n・傷を付けないでください。",
        }).json()
        manual_id = created["id"]
        repaired = client.post(f"/api/manuals/{manual_id}/restructure", headers=headers)
        assert repaired.status_code == 200
        blocks = repaired.json()["document"]["blocks"]
        assert any("30分置いて拭きます。" in block["source_text"] for block in blocks)
        request = next(block for block in blocks if "傷を付けないで" in block["source_text"])
        assert request["kind"] == "request"
        edited = client.put(f"/api/manuals/{manual_id}/blocks/{request['id']}", headers=headers,
                            json={"source_text": request["source_text"], "kind": "other", "tags": ["注意", "清掃"]})
        assert edited.status_code == 200
        assert edited.json()["block"]["tags"] == ["その他", "清掃"]
        assert edited.json()["version"] == 4


def test_review_joins_related_request_and_reassigns_tags():
    from app.ingestion import review_source, structure

    document = structure([(46, "お願い\nたわしでこすらないでください。\n傷付いたり、割れる原因になります。\nとびらパッキン")],
                         "お手入れ", merge_wrapped=True)
    document.blocks[1].tags.append("清掃")
    reviewed = review_source(document)
    assert len(reviewed.blocks) == 2
    assert reviewed.blocks[0].source_text == (
        "お願い\nたわしでこすらないでください。傷付いたり、割れる原因になります。")
    assert reviewed.blocks[0].kind == "request"
    assert reviewed.blocks[0].tags == ["お願い", "清掃"]
    assert reviewed.blocks[1].kind == "other"
    assert reviewed.model_dump() == review_source(reviewed).model_dump()


def test_review_joins_uncertain_fragment_but_keeps_page_boundary():
    from app.ingestion import review_source, structure

    document = structure([(1, "汚れを30分\n置いて拭きます。"), (2, "次の説明")], "お手入れ")
    reviewed = review_source(document)
    assert reviewed.blocks[0].source_text == "汚れを30分置いて拭きます。"
    assert reviewed.blocks[0].tags == ["手順"]
    assert reviewed.blocks[1].source_page == 2


def test_review_does_not_merge_separate_safety_instructions():
    from app.ingestion import review_source, structure

    document = structure([(1, "ぬれた手で電源プラグを抜き差ししない感電の原因になります。\n"
                              "本体のお手入れは電源プラグを抜き、本体が冷めてから行う感電・けがの原因になります。")],
                         "安全")
    reviewed = review_source(document)
    assert len(reviewed.blocks) == 2


def test_generation_reviews_source_before_validation(monkeypatch):
    from app.providers import MockAIProvider

    monkeypatch.setattr("app.main.get_provider", lambda settings, purpose: MockAIProvider())
    with TestClient(app) as client:
        login = client.post("/api/login", json={"username": "admin", "password": "manutailor-demo"}).json()
        headers = {"Authorization": "Bearer " + login["token"]}
        user_id = next(user["id"] for user in client.get("/api/users", headers=headers).json()
                       if user["username"] == "demo")
        created = client.post("/api/manuals", headers=headers, data={
            "title": "検品前の連結", "text": "お願い\nたわしでこすらないでください。\n傷付く原因になります。",
        }).json()
        generated = client.post("/api/generations", headers=headers,
                                json={"manual_id": created["id"], "user_id": user_id})
        assert generated.status_code == 200
        assert generated.json()["report"]["status"] == "pass"
        manual = client.get(f"/api/manuals/{created['id']}", headers=headers).json()
        assert manual["current_version"] == 2
        assert len(manual["document"]["blocks"]) == 1
        result = client.get(f"/api/generations/{generated.json()['id']}", headers=headers).json()
        assert len(result["blocks"]) == 1
        assert result["blocks"][0]["tags"] == ["未整理"]


def test_review_joins_generated_note_and_restores_source_tags():
    from app.ingestion import review_source, structure
    from app.models import GeneratedBlock
    from app.validation import review_generated

    document = review_source(structure([(1, "お願い\n傷を付けないでください。")], "注意"))
    source = document.blocks[0]
    parts = [GeneratedBlock(id="a", generated_text="お願い", source_block_ids=[source.id],
                            reason="layout", tags=["手順"]),
             GeneratedBlock(id="b", generated_text="\n傷を付けないでください。", source_block_ids=[source.id],
                            reason="layout", tags=["手順"])]
    reviewed = review_generated(document, parts)
    assert len(reviewed) == 1
    assert reviewed[0].generated_text == source.source_text
    assert reviewed[0].tags == ["お願い"]


def test_conditional_guidance_is_not_numbered_as_a_step():
    from app.ingestion import structure

    document = structure([(1, "汚れがひどい場合は、台所用洗剤で拭いてください。\n"
                              "汚れを拭き取ってください。")], "お手入れ")
    assert document.blocks[0].kind == "other"
    assert document.blocks[0].tags == ["その他", "条件付き対応"]
    assert document.blocks[1].kind == "step"


def test_manual_merge_and_reclassification_are_versioned():
    with TestClient(app) as client:
        login = client.post("/api/login", json={"username": "admin", "password": "manutailor-demo"}).json()
        headers = {"Authorization": "Bearer " + login["token"]}
        created = client.post("/api/manuals", headers=headers, data={
            "title": "手動連結", "text": "1. 汚れがひどい場合は、台所用洗剤で拭いてください。\n2. 乾いた布で拭き取ります。",
        }).json()
        reparsed = client.post(f"/api/manuals/{created['id']}/restructure", headers=headers).json()
        first, second = reparsed["document"]["blocks"]
        merged = client.post(f"/api/manuals/{created['id']}/blocks/merge", headers=headers, json={
            "first_id": first["id"], "second_id": second["id"],
            "source_text": "汚れがひどい場合は、台所用洗剤で拭いてください。",
        })
        assert merged.status_code == 200
        assert merged.json()["version"] == 4
        assert len(merged.json()["document"]["blocks"]) == 1
        assert merged.json()["document"]["blocks"][0]["tags"] == ["その他", "条件付き対応"]
        edited = client.put(f"/api/manuals/{created['id']}/blocks/{first['id']}", headers=headers,
                            json={"source_text": "汚れがひどい場合は、台所用洗剤で拭いてください。",
                                  "kind": "step", "step_label": "A", "tags": ["清掃"]})
        assert edited.status_code == 200
        assert edited.json()["block"]["step_label"] == "A"
        reclassified = client.post(f"/api/manuals/{created['id']}/reclassify", headers=headers)
        assert reclassified.json()["version"] == 6
        assert reclassified.json()["document"]["blocks"][0]["tags"] == ["その他", "条件付き対応", "清掃"]
        assert reclassified.json()["document"]["blocks"][0]["step_label"] is None


def test_image_tags_and_crop_create_a_new_manual_version():
    image = BytesIO()
    Image.new("RGB", (200, 100), "blue").save(image, format="PNG")
    with TestClient(app) as client:
        login = client.post("/api/login", json={
            "username": "admin", "password": "manutailor-demo",
        }).json()
        headers = {"Authorization": "Bearer " + login["token"]}
        result = client.post(
            "/api/manuals", headers=headers,
            data={"title": "図入り手順", "text": "1. 図を確認してください。"},
            files={"file": ("diagram.png", image.getvalue(), "image/png")},
        )
        assert result.status_code == 200
        manual_id = result.json()["id"]
        figure = result.json()["document"]["images"][0]
        assert figure["tags"] == ["図", "ページ1"]
        edited = client.put(
            f"/api/manuals/{manual_id}/images/{figure['id']}", headers=headers,
            json={"tags": ["操作パネル", "電源"], "crop": {
                "left": 0.25, "top": 0, "right": 0.75, "bottom": 1,
            }},
        )
        assert edited.status_code == 200
        assert edited.json()["version"] == 3
        changed = edited.json()["image"]
        assert changed["tags"] == ["操作パネル", "電源"]
        assert changed["image_path"] != figure["image_path"]
        detail = client.get(f"/api/manuals/{manual_id}", headers=headers).json()
        assert detail["current_version"] == 3
        assert detail["document"]["images"][0]["image_path"] == changed["image_path"]
        from app.config import ROOT
        with Image.open(ROOT / "uploads" / changed["image_path"]) as cropped:
            assert cropped.size == (100, 100)
        target = next(user for user in client.get("/api/users", headers=headers).json()
                      if user["username"] == "demo")
        generated = client.post("/api/generations", headers=headers, json={
            "manual_id": manual_id, "user_id": target["id"],
        })
        assert generated.status_code == 200
        review = client.get(f"/api/generations/{generated.json()['id']}", headers=headers).json()
        assert changed["id"] in review["blocks"][0]["image_ids"]
        assert review["document"]["images"][0]["tags"] == ["操作パネル", "電源"]
