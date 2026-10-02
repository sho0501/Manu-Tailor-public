import io
import re
import subprocess
import shutil
from pathlib import Path

from docx import Document as WordDocument
from PIL import Image
from pypdf import PdfReader

from .config import ROOT
from .db import uid
from .models import Document, ImageBlock, RawPage, SourceBlock, SourceSection

ALLOWED = {".pdf", ".png", ".jpg", ".jpeg", ".txt", ".md", ".docx"}
WARNING = re.compile(r"注意|警告|危険|禁止|必ず|ないで|なりません|するな|不可")
LABELS = {"step": "手順", "warning": "注意", "request": "お願い", "other": "その他"}
ACTION = re.compile(r"(?:してください|して下さい|しないでください|しないで下さい|ます|ましょう|拭く|拭き取る|行う|置く|入れる|外す|押す|回す|確認する|取り出す)[。.!！]?$"
                    r"|(?:てください|て下さい)[。.!！]?$", re.I)
NOTE = re.compile(r"注意|警告|危険|禁止|お願い|故障|けが|傷付|原因にな|おそれ|しないで|ないでください")
CONDITIONAL = re.compile(r"(?:場合|とき|時|際)(?:は|に|、)|(?:なら|たら)(?:、|\s)")
MARKER = re.compile(r"^\s*(?:\d+[.．、)）]|[-*・])\s*")
SENTENCE_END = re.compile(r"[。.!！?？]$|[。.!！?？][」』）)]$")
REASON = re.compile(r"^(?:傷付|さび|故障|破損|感電|やけど|変形|損傷|また、\s*さび)|^.{0,10}(?:原因|おそれ|ことがあり)")
AUTO_TAGS = {*LABELS.values(), "文章要確認", "条件付き対応"}


def classify(text: str, context: str = "") -> tuple[str, list[str]]:
    if context in {"warning", "request"}:
        kind = context
    elif NOTE.search(text):
        kind = "request" if "お願い" in text else "warning"
    elif CONDITIONAL.search(text):
        kind = "other"
    elif ACTION.search(text):
        kind = "step"
    else:
        kind = "other"
    tags = [LABELS[kind]]
    if kind == "other" and CONDITIONAL.search(text):
        tags.append("条件付き対応")
    if kind == "other" and not re.search(r"[。.!！]$", text):
        tags.append("文章要確認")
    return kind, tags


def ocr(path: Path) -> str:
    executable = shutil.which("tesseract")
    if not executable:
        return ""
    result = subprocess.run(
        [executable, str(path), "stdout", "-l", "jpn+eng"], capture_output=True, timeout=60, check=False
    )
    return result.stdout.decode("utf-8", errors="replace").strip() if result.returncode == 0 else ""


def structure(pages: list[tuple[int, str]], title: str, images: list[ImageBlock] | None = None,
              *, merge_wrapped: bool = False) -> Document:
    blocks: list[SourceBlock] = []
    heading = "手順"
    linked_pages: set[int] = set()
    for page, text in pages:
        context = ""
        pending = ""

        def add(value: str, forced_kind: str = "") -> None:
            value = re.sub(r"\s+", " ", value).strip()
            if not value:
                return
            kind, tags = classify(value, forced_kind or context)
            blocks.append(SourceBlock(
                id=uid(), source_page=page, source_block=len(blocks) + 1,
                source_text=value, kind=kind, tags=tags, heading=heading,
                warnings=[value] if kind in {"warning", "request"} or WARNING.search(value) else [],
                image_ids=([im.id for im in (images or []) if im.page == page]
                           if page not in linked_pages else []),
            ))
            linked_pages.add(page)

        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith("#"):
                if pending:
                    add(pending, "other")
                    pending = ""
                heading = line.lstrip("# ")
                context = ""
                continue
            clean = MARKER.sub("", line).strip()
            if clean in {"お願い", "注意", "警告", "危険", "禁止"}:
                if pending:
                    add(pending, "other")
                    pending = ""
                context = "request" if clean == "お願い" else "warning"
                add(clean, context)
                continue
            if not merge_wrapped:
                add(clean)
                continue
            # A short unpunctuated title starts a new section, not a numbered action.
            if (not pending and not MARKER.match(line) and len(clean) <= 25
                    and not re.search(r"[。.!！]", clean)
                    and not re.search(r"(?:は|を|に|で|て|から|場合|\d)", clean)):
                add(clean, "other")
                context = ""
                continue
            if MARKER.match(line) and pending:
                add(pending, "other")
                pending = ""
            pending += clean
            while True:
                match = re.search(r"[。.!！]", pending)
                if not match:
                    break
                sentence, pending = pending[:match.end()], pending[match.end():]
                add(sentence)
                if context in {"warning", "request"} and not pending:
                    context = ""
        if pending:
            add(pending, "other")
    return Document(title=title, blocks=blocks, images=images or [])


def restructure(document: Document) -> Document:
    pages: dict[int, list[str]] = {}
    for block in document.blocks:
        pages.setdefault(block.source_page, []).append(block.source_text)
    updated = structure([(page, "\n".join(lines)) for page, lines in pages.items()],
                        document.title, document.images, merge_wrapped=True)
    updated = review_source(updated)
    updated.source_sections = document.source_sections
    updated.extraction_notes = document.extraction_notes
    updated.raw_pages = document.raw_pages
    updated.visual_groups = document.visual_groups
    updated.excluded_lines = document.excluded_lines
    updated.organization_status = document.organization_status
    return updated


def review_source(document: Document) -> Document:
    """Conservatively join adjacent PDF fragments before generation and validation."""
    result: list[SourceBlock] = []
    for original in document.blocks:
        block = original.model_copy(deep=True)
        previous = result[-1] if result else None
        if previous and previous.source_page == block.source_page:
            before = previous.source_text.strip()
            after = block.source_text.strip()
            standalone_label = before in {"お願い", "注意", "警告", "危険", "禁止"}
            looks_like_fragment = bool(re.search(r"[をはにでて、,]|\d+(?:分|秒|時間|kg|g|ml|L|個|回)", before)) or len(before) > 25
            next_is_heading = (not SENTENCE_END.search(after) and len(after) <= 25
                               and not re.search(r"[をはにでて、,]|\d", after))
            unfinished = (not SENTENCE_END.search(before) and looks_like_fragment
                          and not next_is_heading
                          and after not in {"お願い", "注意", "警告", "危険", "禁止"})
            reason = (previous.kind in {"warning", "request"}
                      and block.kind in {"warning", "request", "other"}
                      and bool(REASON.search(after))
                      and not re.search(r"(?:しないで|してください|して下さい|行う|する)(?:[。.!！]|$)", after))
            same_note = standalone_label and block.kind in {"warning", "request"}
            if (unfinished or reason or same_note) and after:
                separator = "\n" if standalone_label else ""
                previous.source_text = before + separator + after
                context = previous.kind if previous.kind in {"warning", "request"} else ""
                previous.kind, automatic = classify(previous.source_text, context)
                if previous.kind != "step":
                    previous.step_label = None
                custom = [tag for tag in previous.tags + block.tags if tag not in AUTO_TAGS]
                previous.tags = list(dict.fromkeys(automatic + custom))
                previous.warnings = ([previous.source_text] if previous.kind in {"warning", "request"}
                                     or WARNING.search(previous.source_text) else [])
                previous.image_ids = list(dict.fromkeys(previous.image_ids + block.image_ids))
                continue
        result.append(block)
    for index, block in enumerate(result, 1):
        block.source_block = index
    return document.model_copy(update={"blocks": result})


def reclassify_source(document: Document) -> Document:
    blocks = [block.model_copy(deep=True) for block in document.blocks]
    for block in blocks:
        context = block.kind if block.kind in {"warning", "request"} else ""
        kind, automatic = classify(block.source_text, context)
        custom = [tag for tag in block.tags if tag not in AUTO_TAGS]
        block.kind = kind
        if kind != "step":
            block.step_label = None
        block.tags = list(dict.fromkeys(automatic + custom))
        block.warnings = ([block.source_text] if kind in {"warning", "request"}
                          or WARNING.search(block.source_text) else [])
    return document.model_copy(update={"blocks": blocks})


def merge_source_blocks(document: Document, first_id: str, second_id: str, text: str) -> Document:
    blocks = [block.model_copy(deep=True) for block in document.blocks]
    index = next((i for i, block in enumerate(blocks) if block.id == first_id), -1)
    if index < 0 or index + 1 >= len(blocks) or blocks[index + 1].id != second_id:
        raise ValueError("隣り合う項目を選んでください。")
    first, second = blocks[index:index + 2]
    if first.source_page != second.source_page:
        raise ValueError("別のページの項目は連結できません。")
    merged_text = text.strip()
    if not merged_text:
        raise ValueError("連結後の文章を入力してください。")
    first.source_text = merged_text
    first.kind, automatic = classify(merged_text)
    if first.kind != "step":
        first.step_label = None
    custom = [tag for tag in first.tags + second.tags if tag not in AUTO_TAGS]
    first.tags = list(dict.fromkeys(automatic + custom))
    first.warnings = ([merged_text] if first.kind in {"warning", "request"}
                      or WARNING.search(merged_text) else [])
    first.image_ids = list(dict.fromkeys(first.image_ids + second.image_ids))
    del blocks[index + 1]
    for number, block in enumerate(blocks, 1):
        block.source_block = number
    return document.model_copy(update={"blocks": blocks})


def pdf_sections(reader: PdfReader) -> list[SourceSection]:
    """Use bookmark page spans as compact, complete classification units."""
    headings: dict[int, list[str]] = {}

    def visit(items: list) -> None:
        for item in items:
            if isinstance(item, list):
                visit(item)
                continue
            try:
                page = reader.get_destination_page_number(item) + 1
                name = str(item.title).strip()
            except (AttributeError, KeyError, ValueError):
                continue
            if 1 <= page <= len(reader.pages) and name:
                headings.setdefault(page, []).append(name)

    visit(reader.outline)
    if not headings:
        for page, source in enumerate(reader.pages, 1):
            first_line = next((line.strip() for line in (source.extract_text() or "").splitlines()
                               if line.strip()), f"ページ{page}")
            headings[page] = [first_line[:80]]
    else:
        headings.setdefault(1, ["はじめに"])
    starts = sorted(headings)
    return [SourceSection(
        id=f"page:{start}-{(starts[index + 1] - 1) if index + 1 < len(starts) else len(reader.pages)}",
        title=" / ".join(dict.fromkeys(headings[start]))[:120],
        start_page=start,
        end_page=(starts[index + 1] - 1) if index + 1 < len(starts) else len(reader.pages),
    ) for index, start in enumerate(starts)]


def render_pdf_page(path: Path, page: int) -> bytes | None:
    """Preserve vector diagrams that pypdf cannot return as embedded images."""
    executable = shutil.which("pdftoppm")
    if not executable:
        return None
    target = ROOT / "temp" / f"pdf-page-{uid()}"
    result = subprocess.run(
        [executable, "-f", str(page), "-l", str(page), "-scale-to", "1600",
         "-singlefile", "-png", str(path), str(target)],
        capture_output=True, timeout=30, check=False,
    )
    output = target.with_suffix(".png")
    try:
        return output.read_bytes() if result.returncode == 0 and output.exists() else None
    finally:
        output.unlink(missing_ok=True)


def ingest(path: Path, title: str, *, organize: bool = True) -> Document:
    suffix = path.suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("対応形式はPDF、PNG、JPEG、TXT、Markdown、DOCXです。")
    images: list[ImageBlock] = []
    pages: list[tuple[int, str]] = []
    notes: list[str] = []
    sections: list[SourceSection] = []

    def save_image(data: bytes, page: int, *, page_preview: bool = False):
        image_id = uid()
        image_path = ROOT / "uploads" / f"{image_id}.png"
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 40_000_000:
                raise ValueError("画像が大きすぎます。")
            image.convert("RGB").save(image_path)
        images.append(
            ImageBlock(
                id=image_id, image_path=image_path.name, page=page,
                caption="原文ページ" if page_preview else "原文の図",
                source_reference=f"page:{page}:rendered" if page_preview else f"page:{page}",
                tags=(["原文ページ", f"ページ{page}"] if page_preview
                      else ["図", f"ページ{page}"]),
            )
        )
        return image_path

    if suffix == ".pdf":
        reader = PdfReader(path)
        if len(reader.pages) > 100:
            raise ValueError("PDFは100ページ以下にしてください。")
        sections = pdf_sections(reader)
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            image_paths = [save_image(image.data, number) for image in page.images]
            preview = render_pdf_page(path, number)
            if preview:
                image_paths.append(save_image(preview, number, page_preview=True))
            elif not image_paths:
                notes.append(f"ページ{number}: 図の描画を抽出できませんでした。PDF表示で確認してください。")
            if not text.strip():
                text = "\n".join(ocr(image_path) for image_path in image_paths)
                if not text.strip():
                    notes.append(f"ページ{number}: OCR未抽出。画像を確認し、管理者が原文を入力してください。")
            pages.append((number, text))
    elif suffix in {".png", ".jpg", ".jpeg"}:
        image_path = save_image(path.read_bytes(), 1)
        pages = [(1, ocr(image_path))]
        if not pages[0][1]:
            notes.append("画像原文は保持されています。OCR未設定のため原文の手入力が必要です。")
    elif suffix == ".docx":
        word_document = WordDocument(str(path))
        text = "\n".join(p.text for p in word_document.paragraphs)
        for table in word_document.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        for rel in word_document.part.rels.values():
            if "image" in rel.reltype and not rel.is_external:
                save_image(rel.target_part.blob, 1)
        pages = [(1, text)]
    else:
        data = path.read_bytes()
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp932")
        pages = [(1, text)]
    document = (structure(pages, title, images, merge_wrapped=suffix == ".pdf")
                if organize else Document(title=title, blocks=[], images=images, organization_status="raw"))
    if suffix == ".pdf" and organize:
        document = review_source(document)
    document.raw_pages = [RawPage(page=number, text=value) for number, value in pages]
    document.source_sections = sections
    document.extraction_notes = notes
    return document
